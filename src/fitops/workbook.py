"""Normalize the cleaned workbook while preserving source provenance."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl

from .importer import ParsedPeriod, ParsedSetCell, parse_period, parse_set_cell


@dataclass(frozen=True)
class NormalizedWorkoutRow:
    source_tab: str
    source_row: int
    session_label: str
    day_label: str
    exercise: str
    category: str | None
    target_sets: str | None
    target_reps: str | None
    rest: str | None
    period: ParsedPeriod
    set_cells: tuple[ParsedSetCell, ...]
    early_rpe: str | None
    last_set_rpe: str | None
    notes: str | None
    skipped: bool


@dataclass(frozen=True)
class WorkbookContents:
    workouts: tuple[NormalizedWorkoutRow, ...]
    exercise_library: tuple[dict[str, object], ...]
    cleanup_notes: tuple[str, ...]


def _text(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def load_workbook_contents(path: str | Path, year: int | None = None) -> WorkbookContents:
    """Load the known cleaned-workbook tabs into inspectable Python records."""
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    master = workbook["Master Log"]
    workouts: list[NormalizedWorkoutRow] = []
    for source_row, values in enumerate(master.iter_rows(min_row=2, values_only=True), start=2):
        session = _text(values[0]) or ""
        exercise = _text(values[2]) or ""
        skipped = "Skipped this day" in exercise
        workouts.append(
            NormalizedWorkoutRow(
                source_tab=_text(values[16]) or "Master Log",
                source_row=source_row,
                session_label=session,
                day_label=_text(values[1]) or "",
                exercise=exercise,
                category=_text(values[3]),
                target_sets=_text(values[4]),
                target_reps=_text(values[5]),
                rest=_text(values[6]),
                period=parse_period(session, year=year),
                set_cells=tuple(parse_set_cell(value) for value in values[7:11]),
                early_rpe=_text(values[11]),
                last_set_rpe=_text(values[12]),
                notes=_text(values[15]),
                skipped=skipped,
            )
        )

    library = tuple(
        {
            "canonical_exercise": row[0],
            "category": row[1],
            "times_logged": row[2],
            "raw_name_variants": row[3],
        }
        for row in workbook["Exercise Library"].iter_rows(min_row=2, values_only=True)
        if row[0]
    )
    cleanup_notes = tuple(
        str(row[0]).strip()
        for row in workbook["Cleanup Notes"].iter_rows(values_only=True)
        if row[0]
    )
    return WorkbookContents(tuple(workouts), library, cleanup_notes)
